import sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
for ws in wb.worksheets:
    print(f"=== {ws.title!r} rows={ws.max_row} cols={ws.max_column}")
    for i, row in enumerate(ws.iter_rows(max_row=5, values_only=True)):
        print(i, [c for c in row][:60])
